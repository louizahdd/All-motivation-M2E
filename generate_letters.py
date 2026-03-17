from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_letter(filename, recipient_lines, intro_para, closing_para, work_para=None):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    def add_paragraph(text='', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
        return p

    def add_body_paragraph(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.first_line_indent = Cm(1)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'

    # Expéditeur
    for line in ["Louiza Hadid", "06-50-37-56-47", "louizahdd@gmail.com", "3 rue de Ventspils", "56100 Lorient"]:
        add_paragraph(line, space_after=0)

    doc.add_paragraph()

    # Destinataire
    for line in recipient_lines:
        add_paragraph(line, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)

    doc.add_paragraph()

    add_paragraph("Objet : Candidature au Master M2E", bold=True, space_before=6, space_after=12)
    add_paragraph("Madame, Monsieur,", space_after=6)

    default_work_para = (
        "Par ailleurs, mon poste actuel de chargée de recouvrement, bien qu'éloigné du milieu scolaire, "
        "m'a apporté une rigueur organisationnelle et une aisance dans la communication professionnelle "
        "qui seront des atouts pour enseigner. Gérer un portefeuille de comptes, prioriser les actions, "
        "travailler en équipe : autant de compétences transversales que je souhaite mettre au service de "
        "l'éducation."
    )
    paragraphs = [
        intro_para,
        "Mon parcours en licence Sciences de l'éducation à l'Université de Rennes 2 a joué un rôle déterminant dans la construction de ce projet. J'y ai acquis des bases solides en pédagogie, en psychologie du développement et en didactique. Ces trois années m'ont confirmé que c'est dans l'enseignement du premier degré que je veux m'engager durablement, et votre formation représente pour moi l'étape la plus cohérente pour y parvenir.",
        "Ce projet ne s'est pas construit uniquement dans les livres. C'est sur le terrain que ma vocation s'est précisée, étape après étape. D'abord en animation périscolaire, où j'ai découvert qu'un cadre bienveillant compte autant que ce que l'on transmet. Puis en classe de mer, où j'ai compris que c'est la qualité du lien avec les élèves qui rend tout apprentissage possible. Enfin en tutorat auprès d'apprenants étrangers, où j'ai appris à reformuler, différencier et écouter vraiment. Chaque expérience a renforcé la même certitude : c'est dans ce métier que je veux grandir.",
        work_para if work_para else default_work_para,
        closing_para,
    ]

    for para in paragraphs:
        add_body_paragraph(para)

    # Formule de politesse
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(1)
    run = p.add_run("Je me tiens à votre disposition pour un éventuel entretien et vous prie d'agréer, Madame, Monsieur, l'expression de mes respectueuses salutations.")
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

    # Signature
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("Louiza Hadid")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

    doc.save(filename)
    print(f"Created: {filename}")


# ── Paragraphes spécifiques ──────────────────────────────────────────────────

INTRO_INSPE_AIX_MARSEILLE = (
    "C'est une vocation pour l'enseignement, mûrie au fil des années, qui me pousse aujourd'hui à "
    "présenter ma candidature au Master M2E que vous proposez. Ce n'est pas un choix par défaut : "
    "l'INSPE d'Aix-Marseille est un établissement dans lequel je me projette pleinement. La qualité "
    "de la formation, l'articulation entre théorie et pratique professionnelle, et la rigueur de la "
    "préparation au CRPE correspondent aux exigences que je me fixe. C'est dans ce cadre académique "
    "solide que je souhaite préparer le concours et construire mon avenir dans l'enseignement du "
    "premier degré."
)

CLOSING_INSPE_AIX_MARSEILLE = (
    "Je suis profondément attachée à l'idée que l'enseignement ne se réduit pas à la transmission de "
    "savoirs : il s'agit d'accompagner chaque élève dans son développement, avec attention et exigence. "
    "C'est cette conviction qui me pousse à candidater à l'INSPE d'Aix-Marseille, dont la réputation "
    "et les ressources pédagogiques correspondent au niveau d'engagement que je veux apporter à cette "
    "formation. Je suis déterminée à m'investir pleinement pour réussir le CRPE et devenir "
    "professeure des écoles."
)

INTRO_INSPE_LYON = (
    "C'est une vocation pour l'enseignement, mûrie au fil des années, qui me pousse aujourd'hui à "
    "présenter ma candidature au Master M2E que vous proposez. Ce n'est pas un choix par défaut : "
    "l'INSPE de Lyon est un établissement dans lequel je me projette pleinement. Son inscription au "
    "sein de l'Université Claude Bernard – Lyon 1, la richesse de ses équipes de formation et la "
    "qualité de la préparation au CRPE correspondent aux exigences que je me fixe. C'est dans ce "
    "cadre universitaire et professionnel que je souhaite construire mon avenir dans l'enseignement "
    "du premier degré."
)

CLOSING_INSPE_LYON = (
    "Je suis profondément attachée à l'idée que l'enseignement ne se réduit pas à la transmission de "
    "savoirs : il s'agit d'accompagner chaque élève dans son développement, avec attention et exigence. "
    "C'est cette conviction qui me pousse à candidater à l'INSPE de Lyon, dont la renommée et "
    "l'environnement universitaire correspondent au niveau d'engagement que je veux apporter à cette "
    "formation. Je suis déterminée à m'investir pleinement pour réussir le CRPE et devenir "
    "professeure des écoles."
)

INTRO_ISFEC = (
    "C'est une vocation pour l'enseignement, mûrie au fil des années, qui me pousse aujourd'hui à "
    "présenter ma candidature au Master M2E que vous proposez. Ce n'est pas un choix par défaut : "
    "l'ISFEC Bretagne est l'établissement dans lequel je me projette pleinement. La taille humaine "
    "des promotions, l'accompagnement individualisé et l'ancrage dans les valeurs de l'enseignement "
    "catholique correspondent à la vision de l'éducation que je porte depuis plusieurs années. C'est "
    "dans ce cadre exigeant et bienveillant que je souhaite préparer le CRPE et construire mon avenir "
    "dans l'enseignement du premier degré."
)

CLOSING_ISFEC = (
    "Je suis profondément attachée à l'idée que l'enseignement ne se réduit pas à la transmission de "
    "savoirs : il s'agit d'accompagner chaque élève dans son développement, avec attention et exigence. "
    "Les valeurs portées par l'enseignement catholique — le souci de la personne, la bienveillance et "
    "l'engagement au service de tous — rejoignent ce que je veux incarner en tant qu'enseignante. "
    "C'est cette conviction profonde qui me pousse à candidater en priorité à l'ISFEC Bretagne. "
    "Je suis déterminée à m'investir pleinement dans cette formation."
)

INTRO_INSPE_AMIENS = (
    "C'est une vocation pour l'enseignement, mûrie au fil des années, qui me pousse aujourd'hui à "
    "présenter ma candidature au Master M2E parcours distanciel que vous proposez — et c'est mon "
    "premier choix. Ce n'est pas un choix par défaut : votre formation est celle que j'ai identifiée "
    "en priorité, et ce pour plusieurs raisons. La qualité du diplôme national délivré par l'Université "
    "de Picardie Jules Verne, la rigueur de la préparation au CRPE, et la taille de la promotion — "
    "trente étudiants — garantissent un suivi individualisé et un accompagnement humain que je ne "
    "retrouverais nulle part ailleurs dans ce format. C'est dans ce cadre exigeant que je souhaite "
    "construire mon avenir dans l'enseignement du premier degré."
)

WORK_PARA_INSPE_AMIENS = (
    "Mon choix du format distanciel répond à une contrainte familiale concrète et durable. Ma mère "
    "est en arrêt pour maladie professionnelle depuis plusieurs années ; son état nécessite une "
    "présence régulière à ses côtés. Étant fille unique, je suis la seule personne en mesure "
    "d'assurer cet accompagnement. Un déménagement ou des déplacements réguliers vers un campus "
    "distant sont donc incompatibles avec cette réalité. Votre parcours en distanciel est la seule "
    "formation qui me permette de préparer sérieusement le CRPE sans renoncer à mes responsabilités "
    "familiales."
)

CLOSING_INSPE_AMIENS = (
    "Je suis profondément attachée à l'idée que l'enseignement ne se réduit pas à la transmission de "
    "savoirs : il s'agit d'accompagner chaque élève dans son développement, avec attention et exigence. "
    "Cette conviction, je la porte depuis des années, et c'est elle qui me donne la force de construire "
    "ce projet dans des conditions exigeantes. Je suis convaincue que votre formation à distance, loin "
    "d'être un format de facilité, est celui qui demande le plus d'autonomie, d'organisation et de "
    "rigueur — qualités que j'ai développées et que je mettrai entièrement au service de cette "
    "formation. Je suis déterminée à réussir le CRPE et à devenir professeure des écoles."
)

# ── Génération des lettres ────────────────────────────────────────────────────

create_letter(
    "lettre_INSPE_Amiens.docx",
    [
        "INSPE de l'Académie d'Amiens",
        "Université de Picardie Jules Verne",
        "Institut national supérieur du professorat",
        "et de l'éducation – Hauts-de-France",
        "Amiens (80)",
    ],
    INTRO_INSPE_AMIENS,
    CLOSING_INSPE_AMIENS,
    work_para=WORK_PARA_INSPE_AMIENS,
)

create_letter(
    "lettre_INSPE_AixMarseille_Marseille.docx",
    [
        "INSPE d'Aix-Marseille",
        "Institut national supérieur du professorat",
        "et de l'éducation d'Aix-Marseille",
        "Site de Marseille",
        "Marseille (13)",
    ],
    INTRO_INSPE_AIX_MARSEILLE,
    CLOSING_INSPE_AIX_MARSEILLE,
)

create_letter(
    "lettre_INSPE_AixMarseille_Aix.docx",
    [
        "INSPE d'Aix-Marseille",
        "Institut national supérieur du professorat",
        "et de l'éducation d'Aix-Marseille",
        "Site d'Aix-en-Provence",
        "Aix-en-Provence cedex 01 (13)",
    ],
    INTRO_INSPE_AIX_MARSEILLE,
    CLOSING_INSPE_AIX_MARSEILLE,
)

create_letter(
    "lettre_INSPE_Lyon.docx",
    [
        "INSPE de Lyon",
        "Institut national supérieur du professorat",
        "et de l'éducation",
        "Université Claude Bernard - Lyon 1",
        "Lyon (69)",
    ],
    INTRO_INSPE_LYON,
    CLOSING_INSPE_LYON,
)

create_letter(
    "lettre_ISFEC_Bretagne_Rennes.docx",
    [
        "ISFEC Bretagne",
        "UCO – Facultés libres de l'Ouest",
        "Site de Rennes",
        "Rennes (35)",
    ],
    INTRO_ISFEC,
    CLOSING_ISFEC,
)
